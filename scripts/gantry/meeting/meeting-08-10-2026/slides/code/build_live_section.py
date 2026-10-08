"""Insert the live g_aug section after physical slide 50 into a revised copy of the 08-10 deck.

New slides are clones of existing ones (bullets: 51, table: 49, figure: 4), so background, footer, fonts and the
grey/black build convention carry over. The original deck is not modified.
Numbers come from figures/live_resnet_numbers.json (written by fig_live_resnet.py).
"""
__project_origin__ = "added"

import copy
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.table import Table
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SL = os.path.join(HERE, '..')
SRC = os.path.join(SL, '08-10-2026-additional-states-black-box-noise.pptx')
DST = os.path.join(SL, '08-10-2026-additional-states-black-box-noise_live-g.pptx')
FIG = os.path.join(SL, 'figures')
GREY, BLACK = RGBColor(0x6B, 0x6B, 0x6B), RGBColor(0x11, 0x11, 0x11)
FOOT = 'Live start of g_aug: seed 1, low-pass noise data set, best checkpoints'
R_EMBED = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'


def dup(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag.split('}')[1] in ('sp', 'pic', 'graphicFrame', 'grpSp'):
            tree.remove(el)
    rmap = {rid: new.part.relate_to(rel._target, rel.reltype)
            for rid, rel in src.part.rels.items() if rel.reltype.endswith('/image')}
    for el in src.shapes._spTree:
        if el.tag.split('}')[1] in ('sp', 'pic', 'graphicFrame', 'grpSp'):
            c = copy.deepcopy(el)
            for b in c.iter():
                if b.get(R_EMBED) in rmap:
                    b.set(R_EMBED, rmap[b.get(R_EMBED)])
            tree.append(c)
    bg = src._element.cSld.bg
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    # notes: copy the source notes shapes (the default new notes page has no body placeholder here)
    ntree = new.notes_slide.shapes._spTree
    for el in list(ntree):
        if el.tag.split('}')[1] in ('sp', 'pic'):
            ntree.remove(el)
    for el in src.notes_slide.shapes._spTree:
        if el.tag.split('}')[1] in ('sp', 'pic'):
            ntree.append(copy.deepcopy(el))
    return new


def shape(s, name):
    return next(sh for sh in s.shapes if sh.name == name)


def drop(sh):
    sh._element.getparent().remove(sh._element)


def set_text(sh, text, color=None):
    """Replace the text of a text box, keeping the first run's formatting."""
    tf = sh.text_frame
    for extra in tf.paragraphs[1:]:
        tf._txBody.remove(extra._p)
    p = tf.paragraphs[0]
    for r in p.runs[1:]:
        p._p.remove(r._r)
    p.runs[0].text = text
    if color is not None:
        p.runs[0].font.color.rgb = color


def clone_below(s, sh, dy, text, color=None):
    c = copy.deepcopy(sh._element)
    s.shapes._spTree.append(c)
    new = s.shapes[-1]
    new.top = sh.top + dy
    set_text(new, text, color)
    return new


def header(s, num, title, foot=FOOT):
    set_text(shape(s, 'Text 1'), str(num))
    set_text(shape(s, 'Text 2'), foot)
    set_text(shape(s, 'Text 3'), title)


def bullet_slides(prs, base, num, title, bullets, note):
    """One physical slide per build, older bullets grey, the newest black (deck convention)."""
    out = []
    for k in range(1, len(bullets) + 1):
        s = dup(prs, base)
        header(s, num, title)
        for j, name in enumerate(['Text 4', 'Text 5', 'Text 6', 'Text 7']):
            sh = shape(s, name)
            if j < k:
                set_text(sh, bullets[j], BLACK if j == k - 1 else GREY)
            else:
                drop(sh)
        s.notes_slide.notes_text_frame.text = note
        out.append(s)
    return out


def figure_slide(prs, base, num, title, img, side, note):
    s = dup(prs, base)
    header(s, num, title)
    old = shape(s, 'Image 1')                            # same place and size as the slide 3 figure
    s.shapes.add_picture(img, old.left, old.top, old.width, old.height)
    drop(old)
    a, b = shape(s, 'Text 5'), shape(s, 'Text 7')
    step = b.top - a.top
    set_text(a, *side[0])
    set_text(b, *side[1])
    for k, extra in enumerate(side[2:], start=1):
        clone_below(s, b, step * k, *extra)
    s.notes_slide.notes_text_frame.text = note
    return s


def table_slide(prs, base, num, title, rows, caption, foot_lines, note):
    s = dup(prs, base)
    header(s, num, title)
    set_text(shape(s, 'Text 4'), caption)
    frame = shape(s, 'Table 0')
    tbl = frame._element.graphic.graphicData.tbl
    grid = tbl.tblGrid
    total = sum(int(g.get('w')) for g in grid.gridCol_lst)
    ncol = len(rows[0])
    while len(grid.gridCol_lst) < ncol:
        grid.append(copy.deepcopy(grid.gridCol_lst[-1]))
    for tr in tbl.tr_lst:
        while len(tr.tc_lst) < ncol:                       # after the last cell, before the row's extLst
            tr.tc_lst[-1].addnext(copy.deepcopy(tr.tc_lst[-1]))
    widths = [total - 3 * int(total * 0.22)] + [int(total * 0.22)] * (ncol - 1)
    for g, wd in zip(grid.gridCol_lst, widths):
        g.set('w', str(wd))
    body = tbl.tr_lst[1:3]
    while len(tbl.tr_lst) < len(rows):
        tbl.append(copy.deepcopy(body[(len(tbl.tr_lst) - 1) % 2]))
    for tag, base_id in (('colId', 20000), ('rowId', 10000)):     # copied columns/rows need unique ids
        for k, el in enumerate(tbl.iter('{http://schemas.microsoft.com/office/drawing/2014/main}' + tag)):
            el.set('val', str(base_id + k))
    t = Table(tbl, frame)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            p = t.cell(i, j).text_frame.paragraphs[0]
            if p.runs:
                for r in p.runs[1:]:
                    p._p.remove(r._r)
                p.runs[0].text = val
            else:
                p.text = val
    height = sum(r.height for r in t.rows)
    frame.height = height
    bg = shape(s, 'Shape 5')
    bg.height = height + Emu(25400)
    for sh in [sh for sh in s.shapes if sh.name in ('Text 6', 'Shape 7', 'Table 1')]:
        drop(sh)
    cap = shape(s, 'Text 4')
    y0 = bg.top + bg.height + Emu(228600) - cap.top
    for k, line in enumerate(foot_lines):
        clone_below(s, cap, y0 + Emu(400050) * k, line)
    s.notes_slide.notes_text_frame.text = note
    return s


def main():
    T = json.load(open(os.path.join(FIG, 'live_resnet_numbers.json'), encoding='utf-8'))
    prs = Presentation(SRC)
    S = list(prs.slides)
    n0 = len(S)
    bul, tab, fig = S[50], S[48], S[3]                      # physical slides 51, 49, 4

    new = []
    new += bullet_slides(prs, bul, 21, "Jan's question: does PS2 still help when g_aug starts live?", [
        '• Until now: the whole last layer started at zero, so g_aug = 0 and x_a started switched off',
        '• Now g_aug starts nonzero; f_aug still starts at zero, so the start is still the baseline',
        '• Live MLP: the same tanh network as before, only its x_a output weights random',
        '• Live ResNet: network output zero, plus a linear path W z with random x_a rows',
    ], "• Jan: g_aug should not start at zero. He did not prescribe a scale or an architecture\n"
       "• his 2026 paper: NN(z) + W_a z, zero initial NN output, W_a entries U(-1, 1); its parallel simulation "
       "examples used feedforward networks; the public parallel examples I inspected zero the augmentation\n"
       "• the MLP keeps the previous architecture; the ResNet adds a direct path that can represent linear dynamics\n"
       "• exactly at the start the live g still gets no output-loss gradient (f_aug is zero); it becomes trainable "
       "once the coupling learns, checked in the local preflight (nonzero gradients after the first update)\n"
       "• z = [x_b, x_a, u] has 11 entries: 6 physical states, 2 added states, 3 inputs")
    new.append(table_slide(prs, tab, 22, 'The two live starts, next to the original', [
        ['', 'original MLP', 'live MLP', 'live ResNet'],
        ['hidden layers (2 × 16, tanh)', 'random', 'random', 'random'],
        ['output to x_b (f_aug)', 'zero', 'zero', 'zero'],
        ['output to x_a (g_aug)', 'zero', 'random, ±0.25', 'zero'],
        ['linear path W z to x_a', 'none', 'none', 'random, ±0.30'],
        ['linear path W z to x_b', 'none', 'none', 'zero'],
    ], 'Initial values of the learning function, all output biases zero', [
        'Random weights: PyTorch default, uniform ±1/√(number of inputs): 16 → 0.25, 11 → 0.30',
        'Encoder, x_a part: Xavier uniform, ±√(6 / (inputs + outputs)); x_b part: the analytical map',
    ], "• one rule (nn.Linear default, kaiming_uniform_ with a = sqrt(5)) at different fan-in, not two tuned bounds\n"
       "• it replaces the paper's U(-1, 1) after an overflow (next slide); conditioning only, no stability guarantee\n"
       "• within each architecture plain and PS2 get identical initial parameters for the same seed\n"
       "• the architecture comparison also includes the different live paths, it does not isolate architecture"))
    new += bullet_slides(prs, bul, 23, 'Numerical problems so far', [
        '• Paper scale U(−1, 1) for W: overflow after one update (local test, ResNet, seed 1)',
        '• With the PyTorch scale ±0.30: no immediate overflow, but no stability guarantee',
        '• ResNet with PS2: NaN at update 2,889, best weights kept; cause not traced yet',
        '• Live MLP: cluster jobs cancelled at launch, no results yet',
    ], "• overflow: 400-sample local trial, after one optimizer update\n"
       "• the later NaN is not shown to have the same cause; not traced\n"
       "• the final L-BFGS polish of the PS2 run ran out of GPU memory (16,384 windows in one chunk, 11 GB card): "
       "weights unchanged\n"
       "• 'Training done | 200 ep' in the log is the configured budget; the run stopped at about 17 epochs\n"
       "• MLP: launch/cancellation issues, one cancelled during data loading without a Python traceback; cause not "
       "established")
    v = T['val']
    note_val = ("• validation = closed-loop free run over the six full validation records, mean of per-record RMS "
                "(the checkpoint-selection metric), every 1,300 updates\n"
                "• plain: run 161, job 88532, %s m at 15,600 updates, still running when copied, no NaN in the copied log\n"
                "• PS2: run 162, job 88531; phase 1 ended at the 350 cap, unexplained 0.51; phase 2 0.79 to 0.012 in "
                "2,200 steps\n"
                "• PS2 best checkpoint on test: %s m (record E5, 1.2x controller gain, skipped)\n"
                "• one seed, PS2 stopped by the NaN: preliminary" % (v['plain15600'], v['ps2test']))
    for k in (2, 3):
        side = [('Same data, seed and initial values; only PS2 differs', GREY),
                ('Update 2,600: plain %s m, PS2 %s m' % (v['plain2600'], v['ps22600']), BLACK if k == 2 else GREY)]
        if k == 3:
            side.append(('Plain at 15,600: still %s m' % v['plain15600'], BLACK))
        new.append(figure_slide(prs, fig, 24, 'Live ResNet: PS2 still lowers the error',
                                os.path.join(FIG, 'live_resnet_val_curve.png'), side, note_val))
    note_spec = ("• e = y - y_hat, closed-loop free run over the full validation records, Y axis\n"
                 "• Welch ASD: Hann window, 4096 samples (about 1 Hz), 50 % overlap, fs 4 kHz; PSD averaged over the "
                 "four multisine records VA-S1, Y1, P1, L1; records never joined\n"
                 "• band RMS: rfft Parseval per record, mean over the same four records; RMS, not power\n"
                 "• baseline = the model at the start, before training (detuned physical parameters); on slide 3 the "
                 "black line was the baseline at true parameters, other run and data set, so do not compare the two "
                 "figures line by line\n"
                 "• % removed = 1 - band RMS / baseline band RMS (RMS, not power)\n"
                 "• blue band 106 to 230 Hz, red band 230 to 297 Hz (the earlier absorber-region band)\n"
                 + T['note_extra'])
    sa, sb = T['side_plain'], T['side_ps2']
    new.append(figure_slide(prs, fig, 25, 'Prediction error vs frequency, live ResNet',
                            os.path.join(FIG, 'live_resnet_Y_error_asd_a.png'),
                            [(T['side_base'], GREY), (sa, BLACK)], note_spec))
    new.append(figure_slide(prs, fig, 25, 'Prediction error vs frequency, live ResNet',
                            os.path.join(FIG, 'live_resnet_Y_error_asd_b.png'),
                            [(T['side_base'], GREY), (sa, GREY), (sb, BLACK)], note_spec))

    lst = prs.slides._sldIdLst
    ids = list(lst)[n0:]
    for k, el in enumerate(ids):
        lst.remove(el)
        lst.insert(50 + k, el)
    nN = len(ids)

    t1 = prs.slides[0]
    p = shape(t1, 'Text 3').text_frame.paragraphs[0]
    r = p.add_run()
    r.text = 'with a zero and with a live start of g_aug'
    r.font.size = Emu(254000)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    r.font.name = 'Carlito'
    nt = t1.notes_slide.notes_text_frame
    nt.text = nt.text + "\n• part 1: the zero start and its results; part 2: Jan's question, a live start of g_aug"

    slides = list(prs.slides)
    for s in slides[43:50]:
        f = shape(s, 'Text 2')
        set_text(f, 'Zero start of g_aug (MLP). ' + f.text_frame.text)
    lim = slides[50 + nN]
    last = shape(lim, 'Text 7')
    clone_below(lim, last, shape(lim, 'Text 7').top - shape(lim, 'Text 6').top,
                '• Live start: one seed, ResNet only, the PS2 run ended in a NaN; live MLP pending')
    shift = {'21': '26', '22': '27', '23': '28', '24': '29', '25': '30'}
    for s in slides[50 + nN:58 + nN]:
        f = shape(s, 'Text 1')
        set_text(f, shift.get(f.text_frame.text, f.text_frame.text))
    prs.save(DST)
    print('saved', DST, len(prs.slides), 'slides, new', nN)


if __name__ == '__main__':
    main()
